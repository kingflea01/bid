"""Build a structured index for a tender DOCX."""

import json
import re
import time
from pathlib import Path

import docx
from docx.oxml.ns import qn

from core.extract_docx import extract_with_docx
from core.outline import extract_tender_outline, font_size_pt, bold_ratio


def _metadata(text):
    patterns = {
        "project_name": r"项目名称[：:]\s*(.+?)(?:\s{2,}|\n|$)",
        "project_number": r"(?:项目编号|采购编号)[：:]\s*(.+?)(?:\s{2,}|\n|$)",
        "purchaser": r"采\s*购\s*人[：:]\s*(.+?)(?:\s{2,}|\n|$)",
        "agency": r"(?:代理机构|招标机构)[：:]\s*(.+?)(?:\s{2,}|\n|$)",
        "budget": r"(?:项目预算|预算金额|最高限价)[：:\s]*([0-9,，]+)\s*元",
    }
    result = {}
    for key, pattern in patterns.items():
        match = re.search(pattern, text)
        if match:
            value = match.group(1).strip().lstrip("|").strip()
            result[key] = value.replace(",", "").replace("，", "") if key == "budget" else value
    return result


def _table_index(document):
    tables = []
    for index, table in enumerate(document.tables):
        tbl = table._tbl
        rows = tbl.findall(qn("w:tr"))
        grid_cols = tbl.findall(".//" + qn("w:tblGrid") + "/" + qn("w:gridCol"))
        widths = [col.get(qn("w:w")) for col in grid_cols]
        first_row = []
        if rows:
            first_row = [cell.text.strip() for cell in table.rows[0].cells]
        tables.append(
            {
                "index": index,
                "rows": len(rows),
                "grid_cols": len(grid_cols),
                "grid_widths": widths,
                "grid_span": len(tbl.findall(".//" + qn("w:gridSpan"))),
                "v_merge": len(tbl.findall(".//" + qn("w:vMerge"))),
                "first_row": first_row,
            }
        )
    return tables


def build_index(path):
    path = Path(path)
    timing = {}

    start = time.perf_counter()
    document = docx.Document(str(path))
    timing["load_ms"] = round((time.perf_counter() - start) * 1000, 2)

    start = time.perf_counter()
    text, table_count = extract_with_docx(path)
    timing["text_ms"] = round((time.perf_counter() - start) * 1000, 2)

    start = time.perf_counter()
    outline = extract_tender_outline(document)
    timing["outline_ms"] = round((time.perf_counter() - start) * 1000, 2)

    start = time.perf_counter()
    paragraphs = []
    for index, paragraph in enumerate(document.paragraphs):
        paragraph_text = paragraph.text.strip()
        if not paragraph_text:
            continue
        paragraphs.append(
            {
                "index": index,
                "text": paragraph_text,
                "style": paragraph.style.name if paragraph.style else "",
                "bold": round(bold_ratio(paragraph), 2),
                "size_pt": font_size_pt(paragraph),
            }
        )
    tables = _table_index(document)
    timing["index_ms"] = round((time.perf_counter() - start) * 1000, 2)

    metadata = _metadata(text)
    return {
        "source": str(path),
        "table_count": table_count,
        "char_count": len(text),
        "metadata": metadata,
        "chapters": outline["chapters"],
        "format_chapter": outline["format_chapter"],
        "outline": outline["outline"],
        "paragraphs": paragraphs,
        "tables": tables,
        "timing": timing,
        "text": text,
    }


def write_index(path, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    index = build_index(path)
    text = index.pop("text")
    (output_dir / "extracted.txt").write_text(text, encoding="utf-8")
    (output_dir / "tender.json").write_text(
        json.dumps(index, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return index
