"""Assemble bid.docx from structure.json, tender blocks and Cx/Px configs."""

import json
import re
import time
from pathlib import Path

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from builder.clean import (
    clean_paragraph,
    clean_table,
    fill_content_placeholders,
    paragraph_text,
)
from builder.clone import block_elements, clear_body, clone_elements
from builder.config import field, parse_config

PLACEHOLDER_RE = re.compile(r"\[\[(.+?)\]\]")


def set_run_font(run, size=12, bold=False, color=None):
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "宋体"
    r_pr = run._element.get_or_add_rPr()
    fonts = r_pr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        r_pr.append(fonts)
    fonts.set(qn("w:ascii"), "Times New Roman")
    fonts.set(qn("w:hAnsi"), "Times New Roman")
    fonts.set(qn("w:eastAsia"), "宋体")
    if color:
        run.font.color.rgb = color


def add_heading(document, text, level, page_break=False):
    paragraph = document.add_paragraph()
    if page_break:
        paragraph.paragraph_format.page_break_before = True
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text.strip())
    size = 16 if level == 1 else 14 if level == 2 else 12
    set_run_font(run, size=size, bold=True)
    p_pr = paragraph._p.get_or_add_pPr()
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(max(level - 1, 0)))
    p_pr.append(outline)
    return paragraph


def add_rich_text(document, text, align=None, size=12):
    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    for part in re.split(r"(\[\[.+?\]\])", text):
        if not part:
            continue
        run = paragraph.add_run(part)
        set_run_font(run, size=size)
        if part.startswith("[[") and part.endswith("]]"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def add_placeholder(document, text):
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    set_run_font(run, size=12)
    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return paragraph


def add_cover(document, metadata, fields, terminology, package):
    for _ in range(5):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"{terminology}文件")
    set_run_font(run, size=26, bold=True)

    if package and package != "未指定":
        line = document.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = line.add_run(f"（第 {package} 包）")
        set_run_font(run, size=14)

    for _ in range(3):
        document.add_paragraph()

    for label, value in [
        ("项目名称", fields.get("project_name") or "[[项目名称]]"),
        ("项目编号", fields.get("project_number") or "[[项目编号]]"),
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{label}：{value}")
        set_run_font(run, size=14)

    for _ in range(4):
        document.add_paragraph()

    for label, value in [
        ("供应商", fields.get("supplier_name") or "[[供应商名称]]"),
        ("日期", "[[落款日期]]"),
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(f"{label}：{value}")
        set_run_font(run, size=14)

    document.add_page_break()


def add_toc(document):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("目录")
    set_run_font(run, size=16, bold=True)
    paragraph.paragraph_format.space_after = Pt(12)

    field_paragraph = document.add_paragraph()
    run = field_paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在 Word 中更新后显示"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)
    document.add_page_break()


def resolve_content_template(template_path, skill_dir, raw):
    candidates = [
        Path(template_path).parent / raw,
        Path(skill_dir) / "config" / raw,
        Path.cwd() / raw,
    ]
    for path in candidates:
        if path.exists():
            return path
    return None


def record_placeholders(text, chapter, records, seen):
    for value in PLACEHOLDER_RE.findall(text):
        key = (chapter, value)
        if key in seen:
            continue
        seen.add(key)
        records.append({"chapter": chapter, "value": value})


def scan_element_placeholders(element, chapter, records, seen):
    if element.tag == qn("w:p"):
        record_placeholders(paragraph_text(element), chapter, records, seen)
        return
    for cell in element.findall(".//" + qn("w:tc")):
        text = "".join(node.text or "" for node in cell.iter(qn("w:t")))
        record_placeholders(text, chapter, records, seen)


def render_user_content(
    document, node, template_path, skill_dir, chapter, records, seen, fields
):
    mode = node.get("content_mode")
    if mode == "content_template" and node.get("content_template"):
        path = resolve_content_template(
            template_path, skill_dir, node["content_template"]
        )
        if path and path.exists():
            for line in path.read_text(encoding="utf-8").splitlines():
                line = line.strip()
                if not line or line.startswith("#"):
                    continue
                line = fill_content_placeholders(line, fields)
                add_rich_text(document, line, size=12)
                record_placeholders(line, chapter, records, seen)
        else:
            placeholder = f"[[待编写：{node['title']}]]"
            add_placeholder(document, placeholder)
            record_placeholders(placeholder, chapter, records, seen)
        return
    if mode == "fixed_text" and node.get("content_lines"):
        for line in node["content_lines"]:
            line = fill_content_placeholders(line, fields)
            add_rich_text(document, line, size=12)
            record_placeholders(line, chapter, records, seen)
        return
    placeholder = f"[[待编写：{node['title']}]]"
    add_placeholder(document, placeholder)
    record_placeholders(placeholder, chapter, records, seen)


def normalize_title(text):
    value = re.sub(r"^\d+(\.\d+)*[\.、]?\s*", "", text)
    value = re.sub(r"[；。：，、\s]", "", value)
    return value.replace("的", "")


def should_skip(node, extracted_text):
    title = node.get("title", "")
    if "报价说明" in title:
        return True
    if "保证金" in title:
        if (
            "不要求递交" in extracted_text
            or "不再向投标人收取投标保证金" in extracted_text
        ):
            return True
    return False


def render_node(
    document,
    source_document,
    node,
    depth,
    is_first_top,
    fields,
    template_path,
    skill_dir,
    records,
    seen,
    cleanup,
    counters,
    missing,
    extracted_text,
):
    counters["rendered"] += 1
    chapter = node["title"]
    page_break = depth == 1 and not is_first_top

    if should_skip(node, extracted_text):
        cleanup.append(
            {"chapter": chapter, "action": "不适用章节，已省略", "text": chapter}
        )
        return

    if node["source"] == "tender":
        add_heading(document, chapter, min(depth, 3), page_break=page_break)
        elements = block_elements(
            source_document, node.get("source_index"), node.get("source_end")
        )
        if not elements:
            if not node.get("children"):
                counters["missing_blocks"] += 1
                missing.append(chapter)
                placeholder = f"[[待编写：{chapter}]]"
                add_placeholder(document, placeholder)
                record_placeholders(placeholder, chapter, records, seen)
        else:
            cloned = clone_elements(document, elements)
            seen_nonempty = False
            kept = 0
            for element in cloned:
                if element.tag == qn("w:p"):
                    text = paragraph_text(element)
                    if text and not seen_nonempty:
                        seen_nonempty = True
                        if normalize_title(text) == normalize_title(chapter):
                            parent = element.getparent()
                            if parent is not None:
                                parent.remove(element)
                            counters["removed_duplicates"] += 1
                            cleanup.append(
                                {
                                    "chapter": chapter,
                                    "action": "删除重复标题",
                                    "text": text,
                                }
                            )
                            continue
                    action, text = clean_paragraph(element, fields)
                    if action == "remove":
                        counters["removed_paragraphs"] += 1
                        cleanup.append(
                            {"chapter": chapter, "action": "删除说明", "text": text}
                        )
                        continue
                    if action in ("filled", "attach"):
                        cleanup.append(
                            {"chapter": chapter, "action": "替换/填充", "text": text}
                        )
                    if action in ("keep", "filled", "attach"):
                        kept += 1
                    scan_element_placeholders(element, chapter, records, seen)
                elif element.tag == qn("w:tbl"):
                    cleared = clean_table(element)
                    if cleared:
                        counters["cleared_table_rows"] += cleared
                        cleanup.append(
                            {
                                "chapter": chapter,
                                "action": "清空表格数据行",
                                "text": f"清空 {cleared} 行",
                            }
                        )
                    kept += 1
                    scan_element_placeholders(element, chapter, records, seen)

            if kept == 0 and not node.get("children"):
                counters["missing_blocks"] += 1
                missing.append(chapter)
                placeholder = f"[[待编写：{chapter}]]"
                add_placeholder(document, placeholder)
                record_placeholders(placeholder, chapter, records, seen)
    else:
        add_heading(document, chapter, min(depth, 3), page_break=page_break)
        render_user_content(
            document, node, template_path, skill_dir, chapter, records, seen, fields
        )

    for child in node.get("children", []):
        render_node(
            document,
            source_document,
            child,
            depth + 1,
            False,
            fields,
            template_path,
            skill_dir,
            records,
            seen,
            cleanup,
            counters,
            missing,
            extracted_text,
        )


def build_bid(
    output_dir,
    tender_path,
    skill_dir,
    company_code="C1",
    product_code="P1",
):
    started = time.perf_counter()
    output_dir = Path(output_dir)
    skill_dir = Path(skill_dir)
    structure = json.loads((output_dir / "structure.json").read_text(encoding="utf-8"))
    metadata = structure.get("metadata", {})
    template_path = structure.get("template", "")

    company = parse_config(skill_dir / "config" / "companies" / f"{company_code}.md")
    product = parse_config(skill_dir / "config" / "products" / f"{product_code}.md")

    fields = {
        "project_name": metadata.get("project_name") or "[[项目名称]]",
        "project_number": metadata.get("project_number") or "[[项目编号]]",
        "purchaser": metadata.get("purchaser") or "[[采购人]]",
        "supplier_name": field(company, "公司名称", default="[[供应商名称]]"),
        "credit_code": field(company, "统一社会信用代码", default=""),
        "registered_address": field(company, "注册地址", default="[[注册地址]]"),
        "legal_rep": field(company, "法定代表人", default="[[法定代表人]]"),
        "established": field(company, "成立时间", default=""),
        "capital": field(company, "注册资本", default=""),
        "main_business": field(company, "主营业务", default=""),
        "qualifications": "、".join(company.get("特定资质", []) or []),
        "product_name": field(product, "产品名称（脱敏）", "产品名称", default=""),
        "product_version": field(product, "版本 / 年份", "版本", default=""),
    }

    extracted = (output_dir / "extracted.txt").read_text(encoding="utf-8")
    terminology = "响应" if "响应文件" in extracted else "投标"

    source_document = docx.Document(str(tender_path))
    document = docx.Document(str(tender_path))
    clear_body(document)

    add_cover(document, metadata, fields, terminology, metadata.get("package", ""))
    add_toc(document)

    records = []
    seen = set()
    cleanup = []
    counters = {
        "rendered": 0,
        "removed_paragraphs": 0,
        "removed_duplicates": 0,
        "cleared_table_rows": 0,
        "missing_blocks": 0,
    }
    missing = []

    first_top = True
    for node in structure.get("nodes", []):
        render_node(
            document,
            source_document,
            node,
            1,
            first_top,
            fields,
            template_path,
            skill_dir,
            records,
            seen,
            cleanup,
            counters,
            missing,
            extracted,
        )
        first_top = False

    bid_path = output_dir / "bid.docx"
    document.save(str(bid_path))

    reopen = docx.Document(str(bid_path))
    paragraphs = len(reopen.paragraphs)
    tables = len(reopen.tables)

    placeholder_lines = ["# 占位符清单", ""]
    placeholder_lines.append(f"- 项目：{fields['project_name']}")
    placeholder_lines.append(f"- 包号：{metadata.get('package', '未指定')}")
    placeholder_lines.append("")
    placeholder_lines.append("| 章节 | 占位符 | 类型 |")
    placeholder_lines.append("|---|---|---|")
    for record in records:
        value = record["value"]
        if value.startswith("附件："):
            kind = "附件"
        elif value.startswith("待编写："):
            kind = "章节"
        else:
            kind = "字段"
        placeholder_lines.append(f"| {record['chapter']} | [[{value}]] | {kind} |")
    placeholder_lines.append("")
    placeholder_lines.append("## 章节写作提示")
    placeholder_lines.append("")

    def collect_hints(nodes):
        for node in nodes:
            for hint in node.get("hints", []):
                placeholder_lines.append(f"- {node['title']}：{hint}")
            collect_hints(node.get("children", []))

    collect_hints(structure.get("nodes", []))
    (output_dir / "placeholders.md").write_text(
        "\n".join(placeholder_lines), encoding="utf-8"
    )

    cleanup_lines = ["# 模板清洗记录", ""]
    cleanup_lines.append("| 章节 | 操作 | 内容 |")
    cleanup_lines.append("|---|---|---|")
    for item in cleanup:
        text = item["text"].replace("\n", " ").strip()
        if len(text) > 100:
            text = text[:100] + "…"
        cleanup_lines.append(f"| {item['chapter']} | {item['action']} | {text} |")
    (output_dir / "template_cleanup.md").write_text(
        "\n".join(cleanup_lines), encoding="utf-8"
    )

    qa_lines = ["== Bid QA =="]
    qa_lines.append(f"chapters_rendered: {counters['rendered']}")
    qa_lines.append(f"tables_in_bid: {tables}")
    qa_lines.append(f"paragraphs_in_bid: {paragraphs}")
    qa_lines.append(f"placeholders: {len(records)}")
    qa_lines.append(f"removed_paragraphs: {counters['removed_paragraphs']}")
    qa_lines.append(f"removed_duplicates: {counters['removed_duplicates']}")
    qa_lines.append(f"cleared_table_rows: {counters['cleared_table_rows']}")
    qa_lines.append(f"placeholder_chapters: {counters['missing_blocks']}")
    if missing:
        qa_lines.append("placeholder_chapter_list: " + "；".join(missing))
    else:
        qa_lines.append("[PASS] all tender chapters have content or children")
    qa_lines.append(f"output: {bid_path}")
    (output_dir / "qa-bid.txt").write_text("\n".join(qa_lines), encoding="utf-8")

    return {
        "bid": str(bid_path),
        "placeholders": str(output_dir / "placeholders.md"),
        "cleanup": str(output_dir / "template_cleanup.md"),
        "qa": str(output_dir / "qa-bid.txt"),
        "counters": counters,
        "placeholders_count": len(records),
        "tables": tables,
        "paragraphs": paragraphs,
        "elapsed_ms": round((time.perf_counter() - started) * 1000, 2),
    }
