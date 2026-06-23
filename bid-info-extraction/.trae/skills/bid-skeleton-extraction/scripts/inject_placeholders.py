#!/usr/bin/env python3
"""
占位符注入工具
功能：将skeleton.json转换为包含占位符的Word模板文件
Usage: python3 inject_placeholders.py <输出文件夹路径>
Output: <输出文件夹>/template.docx
"""

import os
import sys
import json
from pathlib import Path

# 前置检查：验证 python-docx 是否可用
try:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError as e:
    print(f"Missing dependency: python-docx ({e})", file=sys.stderr)
    print("Install with: pip3 install python-docx", file=sys.stderr)
    sys.exit(1)


def add_paragraph(doc, element):
    """添加段落到文档"""
    style_name = element.get("style", "Normal")
    text = element.get("text", "")
    if element.get("is_heading") and element.get("heading_level"):
        level = element["heading_level"]
        heading_style = f"Heading {level}"
        try:
            para = doc.add_heading(text, level=level)
        except Exception:
            para = doc.add_paragraph(text, style=heading_style)
    else:
        para = doc.add_paragraph(text)
        if style_name and style_name != "Normal":
            try:
                para.style = doc.styles[style_name]
            except KeyError:
                pass
    return para


def add_table(doc, element):
    """添加表格到文档"""
    rows_data = element.get("rows", [])
    if not rows_data:
        return
    col_count = max(len(r) for r in rows_data)
    table = doc.add_table(rows=len(rows_data), cols=col_count)
    table.style = "Table Grid"
    for i, row_data in enumerate(rows_data):
        for j, cell_text in enumerate(row_data):
            if j < col_count:
                cell = table.cell(i, j)
                cell.text = cell_text
    return table


def build_docx(skeleton):
    """根据骨架数据构建Word文档"""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    
    doc = Document()
    
    # 1. 生成封面
    cover = skeleton.get("cover_template", {})
    if cover.get("elements"):
        for elem in cover["elements"]:
            add_paragraph(doc, elem)
        doc.add_paragraph()
    
    # 2. 生成各章节
    for chapter in skeleton.get("chapters", []):
        for elem in chapter.get("elements", []):
            # 检查是否是一级章节标题，如果是则插入分页符
            if (elem.get("type") == "paragraph" and 
                elem.get("is_heading") and 
                elem.get("heading_level") == 1):
                # 添加分页符
                p = doc.add_paragraph()
                run = p.add_run()
                br = OxmlElement('w:br')
                br.set(qn('w:type'), 'page')
                run._element.append(br)
            
            if elem["type"] == "paragraph":
                add_paragraph(doc, elem)
            elif elem["type"] == "table":
                add_table(doc, elem)
        doc.add_paragraph()
    
    return doc


def main():
    # 获取输出文件夹路径
    if len(sys.argv) < 2:
        print("用法: python3 inject_placeholders.py <输出文件夹路径>", file=sys.stderr)
        print("示例: python3 inject_placeholders.py outputs/20260619_项目名称", file=sys.stderr)
        sys.exit(1)
    
    output_dir = Path(sys.argv[1])
    
    # 查找skeleton.json文件
    skeleton_path = output_dir / "skeleton.json"
    
    if not skeleton_path.exists():
        print(f"文件不存在: {skeleton_path}", file=sys.stderr)
        print("请先在该文件夹中创建 skeleton.json", file=sys.stderr)
        sys.exit(1)
    
    with open(skeleton_path, "r", encoding="utf-8") as f:
        skeleton = json.load(f)
    
    # 验证骨架结构
    if not isinstance(skeleton, dict):
        print(f"Invalid skeleton.json: expected object, got {type(skeleton).__name__}", file=sys.stderr)
        sys.exit(1)
    
    chapters = skeleton.get("chapters")
    if chapters is None:
        print("Warning: skeleton.json missing 'chapters' field", file=sys.stderr)
    elif not isinstance(chapters, list):
        print(f"Invalid skeleton.json: 'chapters' must be a list", file=sys.stderr)
        sys.exit(1)
    elif len(chapters) == 0:
        print("Warning: skeleton.json has no chapters (0 chapters)", file=sys.stderr)
    
    # 生成Word文档
    doc = build_docx(skeleton)
    output_path = output_dir / "template.docx"
    doc.save(output_path)
    
    # 统计信息
    chapter_count = len(skeleton.get("chapters", []))
    placeholders = set()
    for ch in skeleton.get("chapters", []):
        placeholders.update(ch.get("placeholders", []))
    cover = skeleton.get("cover_template", {})
    placeholders.update(cover.get("placeholders", []))
    
    print(f"已保存到: {output_path}")
    print(f"统计: 章节 {chapter_count} 个, 占位符 {len(placeholders)} 种")
    print(f"占位符列表: {', '.join(sorted(placeholders))}")


if __name__ == "__main__":
    main()
